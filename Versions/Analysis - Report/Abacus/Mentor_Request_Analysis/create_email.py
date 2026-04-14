from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(11)
font.color.rgb = RGBColor(0, 0, 0)

# Subject line
subj = doc.add_paragraph()
subj_run = subj.add_run('Subject: Re: CoStar Data & Atlanta Introductions – CRE Flood Graph Project')
subj_run.bold = True
subj_run.font.size = Pt(11)

doc.add_paragraph('')

# Greeting
doc.add_paragraph('Hi Mike,')

doc.add_paragraph('')

# Body
doc.add_paragraph(
    'Thank you so much for sharing the CoStar resource and for looking into data options for the '
    'flood graph project. I reviewed their offerings for owners and investors, and I agree — CoStar '
    'would be an excellent comprehensive source for historical CRE development data across markets. '
    'It\'s great to know that option is available as we think about scaling this project.'
)

doc.add_paragraph(
    'For the initial prototype, I\'m taking a lean approach by building with free, publicly available '
    'data first. I\'m currently pulling from WDCEP\'s open data portal for the Washington, D.C. market '
    'and the City of Atlanta\'s open data platform for the Atlanta market. This will allow me to validate '
    'the "flood rising" visualization concept and demonstrate the product\'s potential before we invest '
    'in a paid data source. Once the prototype proves out the concept, we can absolutely scale to CoStar '
    'for richer historical data and broader market coverage.'
)

doc.add_paragraph(
    'I also want to thank you for offering to connect me with your Atlanta contacts. I would welcome '
    'the opportunity to speak with Ernest Kiser at Newmark, Nolan Ford at Benoit, and Elgin Pritchett. '
    'Their perspectives from the Atlanta market would be incredibly valuable as I develop the visualization '
    'and think about how this tool can serve different stakeholders in the CRE space.'
)

doc.add_paragraph(
    'On the deliverables front, the prototype is coming together well. I\'m building three complementary '
    'components: an interactive Python/Streamlit dashboard for the dynamic flood graph visualization, '
    'supporting Excel files with the underlying data models, and a Power BI dashboard for an additional '
    'presentation layer. My goal is to have a working version that clearly demonstrates the concept.'
)

doc.add_paragraph(
    'Would it be helpful for you to review the prototype first before I reach out to Ernest, Nolan, and '
    'Elgin? I want to make sure I\'m putting my best foot forward when connecting with them, and your '
    'feedback would help me refine the product ahead of those conversations.'
)

doc.add_paragraph(
    'Thank you again for your continued guidance and support through Project REAP. I\'m excited about '
    'where this project is heading.'
)

doc.add_paragraph('')

# Closing
doc.add_paragraph('Best regards,')
doc.add_paragraph('Jean Baptiste')

doc.save('/home/ubuntu/Email_Response_to_Mike_CoStar_Atlanta.docx')
print('Done')
